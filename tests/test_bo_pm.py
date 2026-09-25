"""Kiem buoc va v16 -> v17 (`scripts/15_bo_pm.py`).

Buoc nay go ky hieu thong ke va thong nhat mo hinh de xuat la lambda=0.05.
Rui ro lon nhat KHONG phai la go sot, ma la:

  1. Tom tat va Bang 4.8 noi hai con so khac nhau. Do chinh la loi da xay ra o
     v16: Tom tat ghi 28%/33% (lambda=0.05) trong khi Bang 4.8 ghi +21,24%/
     +13,44% (lambda=0.01).
  2. Go qua tay. Khi go trang nhat ky sua doi, ban dau vong lap chay tiep den
     tieu de ke sau va nuot luon HAI BANG TRANG BIA cung logo truong. Test
     `test_ban_nop_van_con_trang_bia` chot lai cai bay do.
"""

from __future__ import annotations

import importlib.util
import re
import statistics as st
import sys
from pathlib import Path

import pytest

pytest.importorskip("docx", reason="can python-docx — chi chay tren VPS")

from docx import Document  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
V16 = ROOT / "docs" / "De_an_thac_si_v16.docx"
RUNS = ROOT / "experiments" / "runs"
OMML = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"

pytestmark = pytest.mark.skipif(
    not V16.is_file() or not RUNS.is_dir(), reason="chua co ban v16")


def load():
    spec = importlib.util.spec_from_file_location(
        "bo_pm", ROOT / "scripts" / "15_bo_pm.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules["bo_pm"] = module
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def mod():
    return load()


@pytest.fixture(scope="module")
def so(mod):
    return mod.doc_so(RUNS)


def _chay(mod, tmp_dir: Path, ten: str, them: list[str]) -> Path:
    out = tmp_dir / ten
    sys.argv = (["15_bo_pm.py", "--source", str(V16), "--runs", str(RUNS),
                 "--out", str(out)] + them)
    mod.main()
    return out


@pytest.fixture(scope="module")
def ra_soat(mod, tmp_path_factory) -> Path:
    return _chay(mod, tmp_path_factory.mktemp("a"), "v17.docx", [])


@pytest.fixture(scope="module")
def ban_nop(mod, tmp_path_factory) -> Path:
    return _chay(mod, tmp_path_factory.mktemp("b"), "v17_nop.docx",
                 ["--no-highlight", "--bo-changelog"])


def doan(path: Path) -> list[str]:
    return [p.text for p in Document(str(path)).paragraphs]


def o_bang(path: Path) -> list[str]:
    return [c.text for t in Document(str(path)).tables
            for r in t.rows for c in r.cells]


# ── 1. Tom tat va Bang 4.8 phai noi cung mot con so ─────────────────────


def bang_48(path: Path):
    doc = Document(str(path))
    for t in doc.tables:
        if t.rows and "λ=0,05" in " ".join(c.text for c in t.rows[0].cells):
            return t
    raise AssertionError("khong thay Bang 4.8 mang nhan λ=0,05")


@pytest.mark.parametrize("cohort,cho_doi", [("original", 28), ("active", 33)])
def test_bang_48_khop_con_so_trong_tom_tat(ban_nop, so, cohort, cho_doi):
    """Tom tat ghi 28% va 33%; Bang 4.8 phai lam tron ve dung hai so do."""
    lgn = st.mean(so[(cohort, "lightgcn", "recall@20")].values())
    bt = st.mean(so[(cohort, "bt_dkgrec_l05", "recall@20")].values())
    assert round((bt / lgn - 1) * 100) == cho_doi

    hang = [r for r in bang_48(ban_nop).rows
            if r.cells[0].text.strip() == cohort
            and r.cells[1].text.strip() == "Recall@20"]
    assert len(hang) == 1
    assert hang[0].cells[3].text.strip() == f"{bt:.6f}"
    assert round(float(hang[0].cells[4].text.strip().rstrip("%"))) == cho_doi


def test_bang_48_khong_con_dung_lambda_001(ban_nop, so):
    """Gia tri cua lambda=0,01 khong duoc con o cot mo hinh de xuat."""
    cot = [r.cells[3].text.strip() for r in bang_48(ban_nop).rows[1:]]
    for cohort in ("original", "active"):
        cu = f"{st.mean(so[(cohort, 'bt_dkgrec', 'recall@20')].values()):.6f}"
        assert cu not in cot, f"con gia tri λ=0,01 ({cu}) trong Bang 4.8"


# ── 2. Ky hieu thong ke da roi khoi than bai ────────────────────────────


@pytest.mark.parametrize("cum", ["±", "Welch", "Ghép cặp", "giá trị p"])
def test_than_bai_sach_ky_hieu_thong_ke(ra_soat, cum):
    """Ban ra soat van giu bang nhat ky (no trich lai) nen chi soat doan van."""
    assert not [t for t in doan(ra_soat) if cum in t], f"con {cum!r} trong doan van"


@pytest.mark.parametrize("cum", ["±", "Welch", "Ghép cặp", "giá trị p"])
def test_ban_nop_sach_tuyet_doi(ban_nop, cum):
    assert cum not in "\n".join(doan(ban_nop) + o_bang(ban_nop))


# ── 3. Go trang nhat ky khong duoc go lan sang trang bia ────────────────


def test_ban_nop_van_con_trang_bia(ban_nop, ra_soat):
    """Bay da sap mot lan: vong lap chay tiep toi tieu de ke sau."""
    van = "\n".join(o_bang(ban_nop))
    assert "BỘ CÔNG THƯƠNG" in van, "da go nham trang bia"
    assert "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in van
    cu, moi = Document(str(ra_soat)), Document(str(ban_nop))
    assert len(moi.tables) == len(cu.tables) - 1, "chi duoc go dung 1 bang"
    assert len(moi.inline_shapes) == len(cu.inline_shapes), "mat hinh"


def test_ban_nop_khong_con_trang_nhat_ky(ban_nop):
    assert not [t for t in doan(ban_nop) if "DANH MỤC SỬA ĐỔI" in t]


def test_ban_ra_soat_van_giu_nhat_ky(ra_soat):
    assert [t for t in doan(ra_soat) if "DANH MỤC SỬA ĐỔI" in t]


# ── 4. Danh so bang sau khi xoa Bang 4.9 ────────────────────────────────


def so_hieu(path: Path, muc_luc: bool) -> list[str]:
    out = []
    for p in Document(str(path)).paragraphs:
        if (p.style.name == "List Bullet") != muc_luc:
            continue
        m = re.match(r"^Bảng (4\.\d+)\.", p.text.strip())
        if m:
            out.append(m.group(1))
    return out


def test_danh_so_lien_tuc_sau_khi_xoa(ra_soat):
    assert so_hieu(ra_soat, False) == [f"4.{i}" for i in range(1, 13)]


def test_muc_luc_khop_than_bai(ra_soat):
    assert so_hieu(ra_soat, True) == so_hieu(ra_soat, False)


def test_khong_mat_cong_thuc_hay_hinh(ra_soat):
    cu, moi = Document(str(V16)), Document(str(ra_soat))
    assert (len(moi.element.body.findall(f".//{OMML}oMath"))
            == len(cu.element.body.findall(f".//{OMML}oMath")))
    assert len(moi.inline_shapes) == len(cu.inline_shapes)
    assert len(moi.tables) == len(cu.tables) - 1
