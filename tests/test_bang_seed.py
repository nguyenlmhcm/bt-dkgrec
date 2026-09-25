"""Kiem buoc va v15 -> v16 (`scripts/14_bang_seed.py`).

Buoc nay lam mot viec khac han cac buoc va truoc: no ghi vao Tom tat mot KHANG
DINH ve so lieu ("lan chay thap nhat cua mo hinh de xuat van cao hon lan chay
tot nhat cua LightGCN"). Rui ro lon nhat vi vay khong phai la dinh dang hong,
ma la tai lieu tuyen bo mot dieu ma so lieu khong con do nua -- neu sau nay
chay them seed, hoac thay du lieu.

Nen test o day gac ba thu:
  1. khang dinh trong Tom tat khop voi metrics.json, va script phai DUNG neu no sai;
  2. so lieu tung seed thuc su co mat trong than bai (v15 khong co);
  3. danh so bang Chuong 4 lien tuc, khong lap, khop giua muc luc va than bai.
"""

from __future__ import annotations

import importlib.util
import re
import sys
from pathlib import Path

import pytest

# Chi VPS moi sinh tai lieu; Colab khong cai python-docx.
pytest.importorskip("docx", reason="can python-docx — chi chay tren VPS")

from docx import Document  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
V15 = ROOT / "docs" / "De_an_thac_si_v15.docx"
RUNS = ROOT / "experiments" / "runs"
OMML = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"

pytestmark = pytest.mark.skipif(
    not V15.is_file() or not RUNS.is_dir(),
    reason="chua co ban v15 hoac thu muc runs")


def load():
    spec = importlib.util.spec_from_file_location(
        "bang_seed", ROOT / "scripts" / "14_bang_seed.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules["bang_seed"] = module
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def mod():
    return load()


@pytest.fixture(scope="module")
def so(mod):
    return mod.doc_tung_seed(RUNS)


@pytest.fixture(scope="module")
def patched(mod, tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("docx") / "v16.docx"
    sys.argv = ["14_bang_seed.py", "--source", str(V15), "--runs", str(RUNS),
                "--out", str(out), "--no-highlight"]
    mod.main()
    return out


def than_bai(path: Path) -> str:
    """Chi cac doan van, khong lay o bang."""
    return "\n".join(p.text for p in Document(str(path)).paragraphs)


def toan_van(path: Path) -> str:
    doc = Document(str(path))
    phan = [p.text for p in doc.paragraphs]
    phan += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join(phan)


# ── 1. Khang dinh trong Tom tat phai dung ───────────────────────────────


def test_ba_seed_du_cho_moi_cau_hinh(so):
    assert len(so) == 12, f"can 6 mo hinh x 2 cohort, co {len(so)}"
    for khoa, gia_tri in so.items():
        assert sorted(gia_tri) == [2020, 2021, 2022], f"{khoa} thieu seed"


@pytest.mark.parametrize("cohort", ["original", "active"])
def test_lan_chay_te_nhat_van_hon_lan_chay_tot_nhat(so, cohort):
    """Chinh la cau duoc viet trong Tom tat. Neu test nay do, phai sua Tom tat."""
    lgn = max(so[(cohort, "lightgcn")].values())
    bt = min(so[(cohort, "bt_dkgrec_l05")].values())
    assert bt > lgn, (
        f"{cohort}: BT thap nhat {bt:.6f} khong hon LightGCN cao nhat {lgn:.6f}")


def test_script_dung_lai_neu_khang_dinh_sai(mod, so):
    """Bay ra du lieu gia trong do LightGCN vuot len; script phai tu choi chay."""
    gia = {k: dict(v) for k, v in so.items()}
    gia[("original", "lightgcn")][2020] = 0.99
    with pytest.raises(SystemExit, match="KHONG cao hon"):
        mod.kiem_tra_khong_chong_lan(gia)


# ── 2. So lieu tung seed phai co mat trong than bai ─────────────────────


def test_bang_47_co_du_muoi_hai_dong(patched):
    doc = Document(str(patched))
    bang = [t for t in doc.tables
            if t.rows and "seed 2020" in " ".join(c.text for c in t.rows[0].cells)]
    assert len(bang) == 1, f"can dung 1 bang theo seed, thay {len(bang)}"
    assert len(bang[0].rows) == 13, "1 dong tieu de + 12 dong du lieu"


@pytest.mark.parametrize("cohort,model", [
    ("original", "lightgcn"), ("original", "bt_dkgrec_l05"),
    ("active", "lightgcn"), ("active", "bt_dkgrec_l05"),
])
def test_moi_gia_tri_tung_seed_deu_xuat_hien(patched, so, cohort, model):
    van = toan_van(patched)
    for seed, gia_tri in so[(cohort, model)].items():
        assert f"{gia_tri:.6f}" in van, (
            f"thieu {cohort}/{model}/seed {seed} = {gia_tri:.6f}")


def test_v15_von_khong_co_so_lieu_nay(so):
    """Ly do buoc va nay ton tai. Neu test do, v15 da co san va co the bo buoc."""
    van = toan_van(V15)
    thieu = [f"{v:.6f}" for v in so[("original", "bt_dkgrec_l05")].values()
             if f"{v:.6f}" not in van]
    assert thieu, "v15 da co so lieu tung seed; xem lai co can buoc 14 khong"


# ── 3. Danh so bang Chuong 4 ────────────────────────────────────────────


def so_hieu(path: Path, trong_danh_muc: bool) -> list[str]:
    doc = Document(str(path))
    out = []
    for p in doc.paragraphs:
        la_muc_luc = p.style.name == "List Bullet"
        if la_muc_luc != trong_danh_muc:
            continue
        m = re.match(r"^Bảng (4\.\d+[a-z]?)\.", p.text.strip())
        if m:
            out.append(m.group(1))
    return out


def test_than_bai_danh_so_lien_tuc(patched):
    got = so_hieu(patched, trong_danh_muc=False)
    assert got == [f"4.{i}" for i in range(1, 14)], got


def test_danh_muc_khop_than_bai(patched):
    assert so_hieu(patched, True) == so_hieu(patched, False)


def test_khong_con_so_hieu_chap_va(patched):
    # Chi doan van: bang danh muc sua doi TRICH LAI so hieu cu ("4.4b -> 4.4"),
    # do la viec cua no, khong phai so hieu con sot.
    assert "Bảng 4.4b" not in than_bai(patched), "4.4b la cach chap va cua v15"


def test_v15_von_bi_trung_so_hieu():
    """Chot lai khiem khuyet cu, de khong ai vo tinh khoi phuc."""
    cu = so_hieu(V15, trong_danh_muc=False)
    assert len(cu) != len(set(cu)), "v15 von dung Bang 4.8 hai lan"
    assert "4.4b" in cu


# ── 4. Tom tat ──────────────────────────────────────────────────────────


def test_tom_tat_khong_con_gia_tri_thap_phan(patched):
    """Hai bai mau cua truong co 0 gia tri thap phan bon chu so trong Tom tat."""
    tom = than_bai(patched).split("TÓM TẮT ĐỀ ÁN THẠC SĨ")[1].split("SUMMARY")[0]
    assert not re.findall(r"0\.\d{4,}", tom), "Tom tat khong duoc chua Recall tho"


def test_tom_tat_bo_cum_dich_may(patched):
    tom = than_bai(patched).split("TÓM TẮT ĐỀ ÁN THẠC SĨ")[1].split("LỜI CAM ĐOAN")[0]
    for cum in ("nguyên mẫu cục bộ", "thông qua Neo4j", "addtocart", "Coverage@K"):
        assert cum not in tom, f"Tom tat khong duoc chua {cum!r}"


def test_khong_mat_cong_thuc_hay_hinh(patched):
    cu, moi = Document(str(V15)), Document(str(patched))
    assert (len(moi.element.body.findall(f".//{OMML}oMath"))
            == len(cu.element.body.findall(f".//{OMML}oMath")))
    assert len(moi.inline_shapes) == len(cu.inline_shapes)
    assert len(moi.tables) == len(cu.tables) + 1


def test_ban_nop_sach_to_nen(patched):
    doc = Document(str(patched))
    con = [r for p in doc.paragraphs for r in p.runs
           if r.font.highlight_color is not None]
    assert not con, f"ban nop con {len(con)} run to nen"
